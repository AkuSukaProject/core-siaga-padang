import sqlite3
import pandas as pd
import openpyxl

print("=" * 70)
print("1. PEMERIKSAAN DI ranah_siaga.db")
print("=" * 70)
conn = sqlite3.connect("ranah_siaga.db")
df_tes = pd.read_sql_query("SELECT * FROM tb_tes", conn)
print(f"Total baris tb_tes: {len(df_tes)}")
tea_matches = df_tes[df_tes['nama_tes'].str.contains('TEA|AKHIR|BUKIT|LAHAN|KETINGGIAN', case=False, na=False)]
print("Cocok kata kunci TEA di tb_tes:")
print(tea_matches[['tes_id', 'nama_tes', 'zona', 'kapasitas', 'lat', 'lon']])

print("\n" + "=" * 70)
print("2. PEMERIKSAAN DI data_tes_padang_lengkap_berkoordinat.csv")
print("=" * 70)
df_csv = pd.read_csv("data_tes_padang_lengkap_berkoordinat.csv")
print(f"Kolom CSV: {list(df_csv.columns)}")
print(f"Total baris CSV: {len(df_csv)}")
for col in df_csv.select_dtypes(include='object').columns:
    m = df_csv[df_csv[col].str.contains('TEA|AKHIR|BUKIT', case=False, na=False)]
    if len(m) > 0:
        print(f"Kolom '{col}' memiliki {len(m)} kecocokan:")
        print(m[['Nomor_TES', 'Sektor', 'Zona', 'Lokasi_TES', 'lat', 'lon'] if 'lat' in m.columns else m.columns[:5]].drop_duplicates())

print("\n" + "=" * 70)
print("3. PEMERIKSAAN DI LAMP 1 EXCEL")
print("=" * 70)
wb = openpyxl.load_workbook("LAMP 1 - TABEL IDENTIFIKASI KETERSEDIAAN DAN KEBUTUHAN FASILITAS EVAKUASI TINGKAT ZONA KOTA PADANG (CETAK).xlsx", data_only=True)
print("Daftar Sheet LAMP 1:", wb.sheetnames)
for sname in wb.sheetnames:
    ws = wb[sname]
    found = []
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(r, c).value or '')
            if any(k in val.upper() for k in ['TEA', 'TEMPAT EVAKUASI AKHIR', 'AKHIR']):
                found.append((r, c, val))
    if found:
        print(f"Sheet '{sname}' ditemukan {len(found)} cell:")
        for r, c, val in found[:10]:
            print(f"  Row {r}, Col {c}: {val}")

print("\n" + "=" * 70)
print("4. PEMERIKSAAN DI LAMP 2 EXCEL")
print("=" * 70)
wb2 = openpyxl.load_workbook("LAMP 2 - TABEL KAJIAN RISIKO BENCANA TSUNAMI KOTA PADANG.xlsx", data_only=True)
print("Daftar Sheet LAMP 2:", wb2.sheetnames)
for sname in wb2.sheetnames:
    ws = wb2[sname]
    found = []
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(r, c).value or '')
            if any(k in val.upper() for k in ['TEA', 'TEMPAT EVAKUASI AKHIR', 'AKHIR']):
                found.append((r, c, val))
    if found:
        print(f"Sheet '{sname}' ditemukan {len(found)} cell:")
        for r, c, val in found[:10]:
            print(f"  Row {r}, Col {c}: {val}")
