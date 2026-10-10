import openpyxl
import pandas as pd

wb = openpyxl.load_workbook("LAMP 1 - TABEL IDENTIFIKASI KETERSEDIAAN DAN KEBUTUHAN FASILITAS EVAKUASI TINGKAT ZONA KOTA PADANG (CETAK).xlsx", data_only=True)
ws = wb['TABEL 2 ']

rows = []
for r in range(5, ws.max_row + 1):
    sektor = ws.cell(r, 1).value
    zona = ws.cell(r, 3).value
    tes_nama = ws.cell(r, 8).value
    tea_no = ws.cell(r, 12).value
    tea_loc = ws.cell(r, 13).value
    tea_cap = ws.cell(r, 14).value
    rute = ws.cell(r, 10).value
    if tes_nama or tea_loc:
        rows.append({
            'row': r, 'sektor': sektor, 'zona': zona, 
            'tes_nama': tes_nama, 'tea_no': tea_no, 'tea_loc': tea_loc, 'tea_cap': tea_cap,
            'rute': rute
        })

df = pd.DataFrame(rows)
df['sektor'] = df['sektor'].ffill()

# Filter Sektor H and I
df_hi = df[df['sektor'].str.contains('SEKTOR H|SEKTOR I', case=False, na=False)]
print("Sample Sektor H & I:")
for idx, r in df_hi.head(30).iterrows():
    print(f"Row {r['row']} | {r['sektor']} | {r['zona']} | TES: {r['tes_nama']} | TEA: {r['tea_no']} - {r['tea_loc']}")
